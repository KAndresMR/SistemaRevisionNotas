from collections import defaultdict
import gc
import io
import os
from pathlib import Path
import shutil
import tempfile
import zipfile

from openpyxl import load_workbook
import pandas as pd
import streamlit as st

from auditor import GradeStatus, StudentMatchStatus
from config import FULL_MAP
from engine import run_audit
from excel_preview import get_preview_legend_html, get_sample_preview_html, worksheet_to_html
from excel_reader import ExcelStructureError
from security import sanitize_filename, validate_excel_file
from teacher_highlighter import create_or_update_discrepancies_sheet, highlight_teacher_file


def guess_course_from_filename(filename: str, fmap: dict) -> str | None:
    norm_name = filename.upper().replace(" ", "").replace("_", "").replace("-", "")
    sorted_keys = sorted(fmap.keys(), key=len, reverse=True)
    for key in sorted_keys:
        if key.upper() in norm_name:
            return key
    return None


def resolve_school_logo(custom_path_str: str | None = None) -> Path | None:
    """
    Localiza el logotipo oficial de la institución educativa.
    Prioriza:
    1. Ruta personalizada si fue especificada por el usuario.
    2. Raíz del repositorio / espacio de trabajo (logoEscuela.png).
    3. Directorio del script o directorio de ejecución actual.
    """
    candidates = []
    if custom_path_str and custom_path_str.strip():
        candidates.append(Path(custom_path_str.strip()))

    script_dir = Path(__file__).resolve().parent
    candidates.extend([
        script_dir.parent.parent / "logoEscuela.png",  # Workspace root
        script_dir.parent / "logoEscuela.png",         # files/
        script_dir / "logoEscuela.png",                # files/src/
        Path.cwd() / "logoEscuela.png",                # CWD
        Path("logoEscuela.png"),
    ])

    for candidate in candidates:
        try:
            if candidate.is_file() and candidate.stat().st_size > 0:
                return candidate.resolve()
        except (OSError, ValueError):
            continue
    return None


st.set_page_config(page_title="Auditoría de Notas - Sistema Institucional", page_icon="📊", layout="wide")

# Diseño Premium para Streamlit
st.markdown("""
<style>
    .main { background-color: #f4f6f9; }
    h1, h2, h3 { color: #1e3a8a; font-family: 'Inter', sans-serif; }
    .stButton>button { background-color: #2563eb; color: white; border-radius: 8px; border: none; padding: 10px 24px; font-weight: bold; transition: 0.3s; }
    .stButton>button:hover { background-color: #1d4ed8; color: white; transform: scale(1.02); }
    .stDownloadButton>button { background-color: #10b981; color: white; border-radius: 8px; border: none; font-weight: bold; }
    .stDownloadButton>button:hover { background-color: #059669; color: white; }
    .excel-preview-card, .excel-preview-card * { color: #0f172a !important; }
    .school-header { display: flex; align-items: center; gap: 18px; margin-bottom: 12px; }
    .school-title { margin: 0; padding: 0; font-size: 2.1rem; font-weight: 800; color: #1e3a8a; line-height: 1.15; }
    .school-subtitle { margin: 4px 0 0 0; color: #475569; font-size: 1.05rem; }
</style>
""", unsafe_allow_html=True)

# Barra lateral institucional
with st.sidebar:
    st.markdown("### 🏫 Estado Institucional")
    auto_logo = resolve_school_logo()
    if auto_logo:
        st.success("✅ **Logo Oficial Detectado**\n\n`logoEscuela.png` se adjunta automáticamente en la esquina superior izquierda de cada Excel exportado.")
    else:
        st.warning("⚠️ No se encontró `logoEscuela.png` en la raíz.")

    with st.expander("⚙️ Opciones avanzadas de ruta", expanded=False):
        custom_logo_input = st.text_input(
            "Ruta personalizada (opcional):",
            value="",
            placeholder="/ruta/hacia/tu/logoEscuela.png",
            help="Solo si moviste el archivo a otra carpeta fuera del proyecto."
        )

    st.markdown("---")
    st.caption("Sistema de Auditoría y Verificación de Calificaciones Académicas.")

# Renderizado de Encabezado Superior con Logotipo a la Izquierda
logo_path = resolve_school_logo(custom_logo_input if 'custom_logo_input' in locals() and custom_logo_input else None)

if logo_path:
    col_logo, col_header = st.columns([1, 6], vertical_alignment="center")
    with col_logo:
        st.image(str(logo_path), width=100)
    with col_header:
        st.markdown('<h1 class="school-title">📊 Dashboard de Revisión de Notas</h1>', unsafe_allow_html=True)
        st.markdown('<p class="school-subtitle"><strong>Escuela de Educación Básica Estados Unidos</strong> • Auditoría de Calificaciones</p>', unsafe_allow_html=True)
else:
    st.title("📊 Dashboard de Revisión de Notas")
    st.markdown("Sube tu matriz de docentes y los reportes oficiales para generar la auditoría automática al instante.")

st.markdown("---")

col1, col2 = st.columns(2)

with col1:
    st.subheader("1. Docente(s)")
    teacher_files = st.file_uploader(
        "Sube uno o varios archivos (ej. notas_8AM.xlsx o Matriz Consolidada)",
        type=['xlsx'],
        accept_multiple_files=True
    )

with col2:
    st.subheader("2. Reportes Oficiales")
    report_files = st.file_uploader(
        "Sube uno o varios reportes (ej. 8AM.xlsx, 9AM.xlsx)",
        type=['xlsx'],
        accept_multiple_files=True
    )

# Guía previa visual con formato idéntico a las capturas oficiales
with st.expander("ℹ️ Vista Previa del Formato Oficial de Exportación (Materias hacia arriba, Docente y Colores)", expanded=False):
    st.markdown('<div class="excel-preview-card">', unsafe_allow_html=True)
    st.markdown(get_preview_legend_html(), unsafe_allow_html=True)
    st.markdown(get_sample_preview_html(), unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)

st.write("---")

# Botón para disparar la ejecución
if st.button("🚀 Ejecutar Auditoría") and teacher_files and report_files:
    # ---------------------------------------------------------
    # VALIDACIÓN DE SEGURIDAD Y FORMATO (5 CAPAS)
    # ---------------------------------------------------------
    for t_file in teacher_files:
        is_valid, error_msg = validate_excel_file(t_file)
        if not is_valid:
            st.error(f"⚠️ Archivo Docente ('{t_file.name}') inválido: {error_msg}")
            st.stop()

    for report in report_files:
        is_valid, error_msg = validate_excel_file(report)
        if not is_valid:
            st.error(f"⚠️ Reporte Oficial ('{report.name}') inválido: {error_msg}")
            st.stop()

    with st.spinner("Procesando y auditando archivos..."):
        # Crear directorio temporal
        temp_dir = Path(tempfile.mkdtemp())

        try:
            # 1. Emparejamiento Inteligente (Smart Match)
            teacher_map = {}
            for t_file in teacher_files:
                c_key = guess_course_from_filename(t_file.name, FULL_MAP)
                if c_key:
                    teacher_map[c_key] = t_file
                else:
                    teacher_map["CONSOLIDATED"] = t_file

            all_discrepancies = []
            file_discrepancies = defaultdict(list)
            file_courses_dict = defaultdict(list)
            course_summary = {}
            processed_courses = []

            progress_bar = st.progress(0)

            # Subdirectorios protegidos
            marked_dir = temp_dir / "resultados_marcados"
            marked_dir.mkdir(exist_ok=True)

            corrected_dir = temp_dir / "resultados_corregidos"
            corrected_dir.mkdir(exist_ok=True)

            teacher_subdir = temp_dir / "docentes"
            teacher_subdir.mkdir(exist_ok=True)

            official_subdir = temp_dir / "oficiales"
            official_subdir.mkdir(exist_ok=True)

            for idx, report in enumerate(report_files):
                course_key = guess_course_from_filename(report.name, FULL_MAP)

                if not course_key:
                    st.warning(f"⚠️ El archivo '{report.name}' tiene un nombre irreconocible. Se omite.")
                    continue

                sheet_name, course_name = FULL_MAP[course_key]

                # 2. Cargas Asimétricas: Buscar pareja
                matched_t_file = teacher_map.get(course_key) or teacher_map.get("CONSOLIDATED")

                if not matched_t_file:
                    st.warning(f"⚠️ Falta archivo de docente para '{course_key}'. Omitiendo este curso.")
                    continue

                # Sanitización de nombres contra Path Traversal / Zip Slip
                safe_teacher_name = sanitize_filename(matched_t_file.name, fallback_prefix=f"docente_{course_key}")
                teacher_path = teacher_subdir / safe_teacher_name
                if not teacher_path.exists():
                    with open(teacher_path, "wb") as f:
                        f.write(matched_t_file.getbuffer())

                safe_report_name = sanitize_filename(report.name, fallback_prefix=f"oficial_{course_key}")
                report_path = official_subdir / safe_report_name
                with open(report_path, "wb") as f:
                    f.write(report.getbuffer())

                marked_teacher_output = marked_dir / f"Matriz_Marcada_{safe_teacher_name}"
                corrected_teacher_output = corrected_dir / f"Matriz_Corregida_{safe_teacher_name}"

                try:
                    marked_path, audit_rows, actual_sheet = run_audit(
                        teacher_file=str(teacher_path),
                        official_file=str(report_path),
                        course_name=course_name,
                        teacher_sheet=sheet_name,
                        teacher_marked_output=str(marked_teacher_output),
                        reset_marked_teacher=False,
                        template_file=str(teacher_path),
                    )

                    highlight_teacher_file(
                        teacher_file=str(teacher_path),
                        output_file=str(corrected_teacher_output),
                        audit_rows=audit_rows,
                        sheet_name=actual_sheet,
                        reset_marked=False,
                        template_file=str(teacher_path),
                        fix_errors=True,
                    )

                    # Métricas
                    n_err = sum(1 for r in audit_rows if r.has_error)
                    n_disc = sum(sum(1 for c in r.comparisons if c.status == GradeStatus.ERROR) for r in audit_rows)
                    n_desf = sum(1 for r in audit_rows if r.has_desfase)

                    course_summary[course_key] = {
                        "total_students": len(audit_rows),
                        "students_with_error": n_err,
                        "total_discrepancies": n_disc,
                        "students_with_desfase": n_desf,
                    }

                    for audit_row in audit_rows:
                        if not audit_row.comparisons:
                            detail = "SOLO EN DOCENTES" if audit_row.match_status == StudentMatchStatus.ONLY_IN_TEACHERS else "SOLO EN SISTEMA"
                            disc = {
                                "curso": course_key,
                                "estudiante": audit_row.student_name,
                                "materia": detail,
                                "docente": "-",
                                "sistema": "-",
                                "diferencia": "-",
                                "estado": "DESFASE DE NÓMINA"
                            }
                            all_discrepancies.append(disc)
                            file_discrepancies[str(marked_teacher_output)].append(disc)
                        else:
                            for comp in audit_row.comparisons:
                                if comp.status == GradeStatus.ERROR:
                                    disc = {
                                        "curso": course_key,
                                        "estudiante": audit_row.student_name,
                                        "materia": comp.subject_display,
                                        "docente": comp.teacher_grade if comp.teacher_grade is not None else "(Vacío)",
                                        "sistema": comp.official_grade if comp.official_grade is not None else "(Vacío)",
                                        "diferencia": comp.difference if comp.difference is not None else "-",
                                        "estado": "ERROR"
                                    }
                                    all_discrepancies.append(disc)
                                    file_discrepancies[str(marked_teacher_output)].append(disc)

                    processed_courses.append(course_key)
                    file_courses_dict[str(marked_teacher_output)].append(course_key)

                except ExcelStructureError as e:
                    st.error(
                        f"⚠️ **Problema estructural en el curso {course_key}**\n\n"
                        f"No se pudo encontrar una estructura de datos válida en el Excel de este curso.\n\n"
                        f"**Detalle técnico:** `{str(e)}`\n\n"
                        f"**💡 Qué revisar:**\n"
                        f"- Asegúrate de que el Excel original tenga las columnas clave como 'ESTUDIANTES APELLIDOS/NOMBRES' o 'NÓMINA'.\n"
                        f"- Revisa que el Excel no esté protegido con contraseña o dañado.\n"
                        f"- Asegúrate de que la hoja de trabajo activa no esté totalmente en blanco.",
                        icon="🚨"
                    )
                except AttributeError as e:
                    import traceback
                    trace = traceback.format_exc()
                    if "MergedCell" in str(e):
                        st.error(
                            f"⚠️ **Problema de Celdas Combinadas en el curso {course_key}**\n\n"
                            f"El sistema intentó escribir información o limpiar colores sobre una celda que está 'Combinada' (Merged) con otras, lo cual bloquea el proceso.\n\n"
                            f"**💡 Qué revisar:**\n"
                            f"- Abre el archivo de Excel del docente para el curso {course_key}.\n"
                            f"- Revisa si en la zona de las notas, promedios o encabezados hay celdas que estén 'Combinadas y Centradas' de forma inusual.\n"
                            f"- Puedes descombinarlas en Excel y volver a subir el archivo.\n\n"
                            f"**Detalle técnico:** `{str(e)}`",
                            icon="🚨"
                        )
                    else:
                        st.error(
                            f"⚠️ **Error de atributos procesando {course_key}**\n\n"
                            f"Ocurrió un problema técnico inesperado al manipular el archivo. Es posible que tenga un formato no soportado.\n\n"
                            f"**💡 Qué revisar:**\n"
                            f"- Comprueba si el formato del archivo es un `.xlsx` estándar y guárdalo de nuevo.\n\n"
                            f"**Detalle técnico:** `{str(e)}`",
                            icon="🚨"
                        )
                        with st.expander("Ver rastro del error completo (para soporte técnico)"):
                            st.code(trace)
                except Exception as e:
                    import traceback
                    trace = traceback.format_exc()
                    st.error(
                        f"🛑 **Error inesperado procesando el curso {course_key}**\n\n"
                        f"El sistema se encontró con un fallo imprevisto durante el procesamiento de este documento.\n\n"
                        f"**💡 Qué revisar:**\n"
                        f"- Asegúrate de que el archivo del docente corresponda al curso y no esté corrupto.\n"
                        f"- Si el problema persiste, intenta abrir el archivo en Excel, 'Guardar como' un archivo nuevo, y volver a subirlo.\n\n"
                        f"**Detalle técnico:** `{str(e)}`",
                        icon="🛑"
                    )
                    with st.expander("Ver rastro del error completo (para soporte técnico)"):
                        st.code(trace)

                progress_bar.progress((idx + 1) / len(report_files))

            # Post-proceso: Insertar hoja de discrepancias dentro de cada archivo marcado
            if file_courses_dict:
                for f_path, courses in file_courses_dict.items():
                    discrepancies = file_discrepancies.get(f_path, [])
                    create_or_update_discrepancies_sheet(f_path, discrepancies, courses)

            # Generar previsualizaciones HTML fieles del archivo Excel generado
            excel_previews = {}
            for f in sorted(marked_dir.glob("*.xlsx")):
                try:
                    wb_prev = load_workbook(f, data_only=True)
                    excel_previews[f.name] = {}
                    for s_name in wb_prev.sheetnames:
                        excel_previews[f.name][s_name] = worksheet_to_html(wb_prev[s_name], max_rows=45, max_cols=18)
                    wb_prev.close()
                except Exception as prev_err:
                    excel_previews[f.name] = {"Error": f"<p style='color:red;'>No se pudo generar vista previa: {str(prev_err)}</p>"}

            # Empaquetar resultados en ZIPs en memoria
            marked_zip_buffer = io.BytesIO()
            with zipfile.ZipFile(marked_zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                for f in marked_dir.glob("*.xlsx"):
                    zip_file.write(f, arcname=f.name)

            corrected_zip_buffer = io.BytesIO()
            with zipfile.ZipFile(corrected_zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                for f in corrected_dir.glob("*.xlsx"):
                    zip_file.write(f, arcname=f.name)

            # Guardar en sesión para que persista ante interacciones y descargas
            st.session_state["audit_completed"] = True
            st.session_state["excel_previews"] = excel_previews
            st.session_state["download_data_marked"] = marked_zip_buffer.getvalue()
            st.session_state["download_data_corrected"] = corrected_zip_buffer.getvalue()
            st.session_state["all_discrepancies"] = all_discrepancies

        finally:
            # LIMPIEZA EXTREMA: Borrar todos los archivos físicos temporales inmediatamente
            shutil.rmtree(temp_dir, ignore_errors=True)
            # Limpiar RAM
            gc.collect()

# =============================================================================
# RENDERIZADO DE RESULTADOS (PERSISTENTE FUERA DEL BLOQUE DEL BOTÓN)
# =============================================================================
if st.session_state.get("audit_completed"):
    dl_marked = st.session_state.get("download_data_marked")
    dl_corrected = st.session_state.get("download_data_corrected")
    discrepancies_list = st.session_state.get("all_discrepancies", [])
    previews = st.session_state.get("excel_previews", {})

    st.success("✅ ¡Auditoría completada exitosamente!")

    if dl_marked and dl_corrected:
        dl_col1, dl_col2 = st.columns(2)
        with dl_col1:
            st.download_button(
                label="📥 Descargar Errores Resaltados (ZIP)",
                data=dl_marked,
                file_name="Resultados_Errores_Marcados.zip",
                mime="application/zip",
                use_container_width=True
            )
        with dl_col2:
            st.download_button(
                label="📥 Descargar Matrices Corregidas (ZIP)",
                data=dl_corrected,
                file_name="Resultados_Matrices_Corregidas.zip",
                mime="application/zip",
                use_container_width=True
            )

    # -------------------------------------------------------------
    # VISTA PREVIA INTERACTIVA DEL EXCEL CON COLORES Y FORMATO REAL
    # -------------------------------------------------------------
    st.write("---")
    st.subheader("👁️ Vista Previa en Vivo del Excel Exportado")
    st.markdown(
        "Inspecciona la apariencia real del archivo Excel generado tal como se exporta: "
        "con colores de nómina, celdas de calificación con error, notas emergentes 💬 y cabeceras institucionales."
    )
    st.markdown('<div class="excel-preview-card">', unsafe_allow_html=True)
    st.markdown(get_preview_legend_html(), unsafe_allow_html=True)

    if previews:
        if len(previews) > 1:
            selected_filename = st.selectbox(
                "📁 Selecciona el archivo para previsualizar:",
                options=list(previews.keys()),
                key="preview_file_selector"
            )
        else:
            selected_filename = list(previews.keys())[0]

        file_sheets = previews.get(selected_filename, {})
        if file_sheets:
            sheet_names = list(file_sheets.keys())
            tabs = st.tabs([f"📄 {sn}" for sn in sheet_names])
            for tab, sn in zip(tabs, sheet_names):
                with tab:
                    st.caption(f"Visualizando hoja **{sn}** de **{selected_filename}** (muestra fiel con formato real):")
                    st.markdown(file_sheets[sn], unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)

    # Mostrar resumen tabular de errores
    st.write("---")
    st.subheader("📊 Resumen Detallado de Errores")
    if discrepancies_list:
        df_errores = pd.DataFrame(discrepancies_list)
        # Convertir columnas a string puro para evitar fallos de PyArrow
        for col in ["docente", "sistema", "diferencia"]:
            if col in df_errores.columns:
                df_errores[col] = df_errores[col].astype(str)

        st.dataframe(df_errores, use_container_width=True)

        errores_por_curso = df_errores['curso'].value_counts()
        st.bar_chart(errores_por_curso)
    else:
        st.info("¡Excelente! No se encontraron discrepancias en los cursos procesados.")
